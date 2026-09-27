import os, json, warnings
import numpy as np
import pandas as pd
import joblib
import openpyxl
from flask import Flask, render_template, request, jsonify
from sklearn.linear_model import LinearRegression

warnings.filterwarnings('ignore')
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"

app = Flask(__name__)

SAVE_DIR  = 'saved_models'
DATA_PATH = os.path.join('dataset', 'Landfilled Waste Composition Dataset.xlsx')
AGG_CACHE = os.path.join(SAVE_DIR, 'agg_cache.csv')
# Set MODEL=lstm to serve a PyTorch LSTM checkpoint instead of the best model
# chosen in metrics_summary.json (requires torch and saved_models/best_model_lstm.pth).
USE_LSTM  = os.environ.get('MODEL', '').lower() == 'lstm'


def make_lstm_class():
    """Defined lazily so torch is only needed when MODEL=lstm."""
    import torch.nn as nn

    class LSTMForecaster(nn.Module):
        def __init__(self, input_size=1, hidden_size=128, num_layers=2, dropout=0.2):
            super().__init__()
            self.lstm = nn.LSTM(input_size, hidden_size, num_layers,
                                batch_first=True, dropout=dropout)
            self.head = nn.Sequential(
                nn.Linear(hidden_size, 64), nn.ReLU(),
                nn.Dropout(0.1), nn.Linear(64, 1)
            )

        def forward(self, x):
            out, _ = self.lstm(x)
            return self.head(out[:, -1, :])

    return LSTMForecaster


def build_df_agg(encoders):
    if os.path.exists(AGG_CACHE):
        print("  Loading cached state-year aggregates ...")
        df = pd.read_csv(AGG_CACHE)
        df['State_enc'] = df['State_enc'].astype(int)
        return df

    print("  Processing dataset (one-time, may take a minute) ...")
    wb = openpyxl.load_workbook(DATA_PATH, data_only=True)
    ws = wb['Dataset']
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows    = list(ws.iter_rows(min_row=2, values_only=True))
    df_raw  = pd.DataFrame(rows, columns=headers)
    df_raw  = df_raw.loc[:, df_raw.columns.notna()]
    df_raw.columns = df_raw.columns.str.strip()

    df = df_raw.copy()
    for col in ['ORD-Designated Facility Type', 'ORD-Designated Waste Stream']:
        if col in df.columns:
            df = df[~df[col].astype(str).str.startswith('=')]
    df['Waste Mass (short tons)'] = pd.to_numeric(df['Waste Mass (short tons)'], errors='coerce')
    df['Year'] = pd.to_numeric(df['Year'], errors='coerce')
    df = df.dropna(subset=['Waste Mass (short tons)', 'Year', 'State'])
    df['Year'] = df['Year'].astype(int)
    df = df[df['Waste Mass (short tons)'] > 0].reset_index(drop=True)

    state_le = encoders['State']
    df = df[df['State'].isin(state_le.classes_)]
    df['State_enc'] = state_le.transform(df['State']).astype(int)

    df_agg = (df.groupby(['State', 'State_enc', 'Year'])['Waste Mass (short tons)']
              .sum().reset_index()
              .sort_values(['State_enc', 'Year'])
              .reset_index(drop=True))
    df_agg['log_total'] = np.log1p(df_agg['Waste Mass (short tons)'])
    df_agg.to_csv(AGG_CACHE, index=False)
    return df_agg


def load_all():
    encoders = joblib.load(os.path.join(SAVE_DIR, 'encoders.joblib'))
    scaler_X = joblib.load(os.path.join(SAVE_DIR, 'scaler_X.joblib'))
    scaler_y = joblib.load(os.path.join(SAVE_DIR, 'scaler_y.joblib'))

    with open(os.path.join(SAVE_DIR, 'metrics_summary.json')) as f:
        raw_metrics = json.load(f)

    # Deduplicate and sort
    seen, unique = set(), []
    for m in raw_metrics:
        if m['Model'] not in seen:
            seen.add(m['Model'])
            unique.append(m)
    metrics_df = pd.DataFrame(unique).sort_values('R2', ascending=False).reset_index(drop=True)

    lstm_path   = os.path.join(SAVE_DIR, 'best_model_lstm.pth')
    sklearn_path = os.path.join(SAVE_DIR, 'best_model.joblib')

    sc_Xl = sc_yl = seq_len = None
    # Serve the best model from metrics_summary.json (saved as best_model.joblib).
    # Previously the app loaded best_model_lstm.pth whenever it existed, so a stale
    # LSTM checkpoint was served while the UI reported Linear Regression.
    if USE_LSTM and os.path.exists(lstm_path):
        import torch
        ckpt    = torch.load(lstm_path, map_location='cpu', weights_only=False)
        model   = make_lstm_class()(**ckpt['arch'])
        model.load_state_dict(ckpt['model_state_dict'])
        model.eval()
        model_type = 'lstm'
        served_name = 'LSTM (PyTorch)'
        sc_Xl   = ckpt['scaler_X']
        sc_yl   = ckpt['scaler_y']
        seq_len = ckpt['seq_len']
    else:
        model      = joblib.load(sklearn_path)
        model_type = 'sklearn'
        served_name = metrics_df.iloc[0]['Model']

    df_agg = build_df_agg(encoders)
    states = sorted(df_agg['State'].unique().tolist())

    ctx = dict(
        encoders=encoders, scaler_X=scaler_X, scaler_y=scaler_y,
        model=model, model_type=model_type,
        sc_Xl=sc_Xl, sc_yl=sc_yl, seq_len=seq_len,
        df_agg=df_agg, states=states,
        metrics=metrics_df.to_dict('records'),
        best_model_name=served_name,
        best_r2=round(float(metrics_df.iloc[0]['R2']), 4),
        best_rmse=round(float(metrics_df.iloc[0]['RMSE']), 2),
        best_mae=round(float(metrics_df.iloc[0]['MAE']), 2),
    )

    # Pre-compute samples for top 6 states by total historical waste
    top6 = (df_agg.groupby('State')['Waste Mass (short tons)']
            .sum().nlargest(6).index.tolist())
    samples = []
    for state in top6:
        try:
            pred = _predict(state, 2026, ctx)
            hist = df_agg[df_agg['State'] == state].sort_values('Year')
            last = hist.iloc[-1] if len(hist) else None
            if len(hist) >= 2:
                prev = float(hist.iloc[-2]['Waste Mass (short tons)'])
                curr = float(hist.iloc[-1]['Waste Mass (short tons)'])
                trend = 'up' if curr >= prev else 'down'
            else:
                trend = 'neutral'
            samples.append({
                'state': state,
                'year': 2026,
                'prediction': int(round(pred)),
                'last_known_year': int(last['Year']) if last is not None else None,
                'last_known_val': int(round(float(last['Waste Mass (short tons)']))) if last is not None else None,
                'trend': trend,
                'data_points': len(hist),
            })
        except Exception:
            pass
    ctx['samples'] = samples
    return ctx


def _predict(state, year, ctx):
    df_agg   = ctx['df_agg']
    enc_val  = int(ctx['encoders']['State'].transform([state])[0])
    state_df = df_agg[df_agg['State_enc'] == enc_val].sort_values('Year')

    if ctx['model_type'] == 'lstm':
        sl     = ctx['seq_len']
        recent = state_df[state_df['Year'] < year].tail(sl)
        vals   = recent['Waste Mass (short tons)'].values.astype(np.float32)
        if len(vals) == 0:
            vals = np.array([1000.0] * sl, dtype=np.float32)
        elif len(vals) < sl:
            vals = np.pad(vals, (sl - len(vals), 0), mode='edge')
        X_s = ctx['sc_Xl'].transform(vals.reshape(1, -1))
        import torch
        inp = torch.FloatTensor(X_s).unsqueeze(-1)
        with torch.no_grad():
            pred_s = ctx['model'](inp).numpy().ravel()[0]
        return float(max(ctx['sc_yl'].inverse_transform([[pred_s]])[0][0], 0))

    recent   = state_df[state_df['Year'] < year].tail(3)
    log_vals = recent['log_total'].values if len(recent) > 0 else np.array([0.0])
    lag1  = float(log_vals[-1]) if len(log_vals) >= 1 else 0.0
    lag2  = float(log_vals[-2]) if len(log_vals) >= 2 else lag1
    lag3  = float(log_vals[-3]) if len(log_vals) >= 3 else lag2
    roll3 = float(np.mean(log_vals[-3:]))
    min_yr    = int(state_df['Year'].min()) if len(state_df) > 0 else year - 20
    yrs_since = year - min_yr

    X   = np.array([[enc_val, year, yrs_since, lag1, lag2, lag3, roll3]], dtype=np.float32)
    X_s = ctx['scaler_X'].transform(X)

    if isinstance(ctx['model'], LinearRegression):
        pred_log = ctx['scaler_y'].inverse_transform([[ctx['model'].predict(X_s)[0]]])[0][0]
    else:
        pred_log = float(ctx['model'].predict(X_s)[0])

    return float(max(np.expm1(pred_log), 0))


# ── Startup ──────────────────────────────────────────────────────────────────
print("\n---  WasteSight AI  ---")
CTX = load_all()
print(f"  Best model : {CTX['best_model_name']}  |  R2 = {CTX['best_r2']}  |  Ready\n")


# ── Routes ────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html',
        states=CTX['states'],
        samples=CTX['samples'],
        metrics=CTX['metrics'],
        best_model=CTX['best_model_name'],
        best_r2=CTX['best_r2'],
        best_rmse=int(CTX['best_rmse']),
        best_mae=int(CTX['best_mae']),
    )


@app.route('/predict', methods=['POST'])
def predict_api():
    body  = request.get_json(force=True) or {}
    state = str(body.get('state', '')).strip()
    try:
        year = int(body.get('year', 2026))
    except (ValueError, TypeError):
        return jsonify(error='Invalid year value'), 400

    if state not in CTX['states']:
        return jsonify(error=f'State not recognised: {state}'), 400
    if not (1980 <= year <= 2035):
        return jsonify(error='Year must be between 1980 and 2035'), 400

    prediction = _predict(state, year, CTX)

    hist = CTX['df_agg'][CTX['df_agg']['State'] == state].sort_values('Year')
    return jsonify(
        state=state,
        year=year,
        prediction=round(prediction, 2),
        model=CTX['best_model_name'],
        hist_years=hist['Year'].tolist(),
        hist_vals=[round(v, 2) for v in hist['Waste Mass (short tons)'].tolist()],
        data_points=len(hist),
    )


if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5000)
